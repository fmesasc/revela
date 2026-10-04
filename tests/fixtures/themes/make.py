#!/usr/bin/env python3
"""Makes the theme fixtures of tests/suites/io.js (run it again only to change them).

Needs python-pptx (pip install python-pptx); LibreOffice (soffice) for the two
LibreOffice ones, from the Impress template Midnightblue it ships.

- office.pptx       python-pptx's default template: the Office theme (Calibri / Calibri),
                    a shape in "Accent 1, darker 25 %" and text in "Accent 2, lighter 40 %".
- noche.pptx        a dark theme "Noche": its own colour scheme, Cambria / Segoe UI, the
                    master's colour map swapped (bg1 = dk1), the master background a theme
                    background style (p:bgRef 1003: a radial gradient), «Section Header»
                    with a background of its own and «Title Only» light (clrMapOvr).
- faceta.potx       a template without slides: theme "Faceta", Trebuchet MS.
- brisa.thmx        an Office theme file alone: "Brisa", Century Gothic.
- google.pptx       like a Google Slides download: theme "Streamline", Montserrat / Lato,
                    layouts with Google's internal names (TITLE, TITLE_AND_BODY…).
- midnightblue.pptx / midnightblue.odp   LibreOffice's Midnightblue template, converted.
"""
import io, os, re, shutil, subprocess, tempfile, zipfile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DROP = ('docProps/thumbnail.jpeg', 'ppt/printerSettings/printerSettings1.bin')


def build(fill_slides):
    prs = Presentation()
    fill_slides(prs)
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


def patch(data, out, edits={}, extra={}, drop_slides=False):
    """Rewrites a .pptx: edits = {part: fn(xml) -> xml}, extra = {part: xml}."""
    src = zipfile.ZipFile(io.BytesIO(data))
    with zipfile.ZipFile(os.path.join(HERE, out), 'w', zipfile.ZIP_DEFLATED) as z:
        for n in src.namelist():
            if n in DROP or (drop_slides and n.startswith('ppt/slides/')):
                continue
            x = src.read(n)
            if n.endswith('.xml') or n.endswith('.rels'):
                s = x.decode('utf-8')
                s = s.replace('<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.presentationml.printerSettings"/>', '')
                s = re.sub(r'<Relationship [^>]*(printerSettings|thumbnail)[^>]*/>', '', s)
                s = re.sub(r'<Override PartName="/docProps/thumbnail.jpeg"[^>]*/>', '', s)
                if n in edits: s = edits[n](s)
                x = s.encode('utf-8')
            z.writestr(n, x)
        for n, s in extra.items():
            z.writestr(n, s)


def scheme(name, c):
    k = ['dk1', 'lt1', 'dk2', 'lt2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    return f'<a:clrScheme name="{name}">' + ''.join(f'<a:{s}><a:srgbClr val="{v}"/></a:{s}>' for s, v in zip(k, c)) + '</a:clrScheme>'


def retheme(xml, name, colors, major, minor, cname=None):
    xml = re.sub(r'<a:theme ([^>]*)name="[^"]*"', lambda m: f'<a:theme {m.group(1)}name="{name}"', xml, count=1)
    xml = re.sub(r'<a:clrScheme .*?</a:clrScheme>', scheme(cname or name, colors), xml, flags=re.S)
    xml = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*', lambda m: m.group(1) + major, xml)
    xml = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]*', lambda m: m.group(1) + minor, xml)
    return xml.replace('<a:fontScheme name="Office">', f'<a:fontScheme name="{name}">')


# ---- office.pptx --------------------------------------------------------------
def office(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = 'Tema de Office'; s.placeholders[1].text = 'Detectado al importar'
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = 'Colores del tema'
    s.placeholders[1].text = 'Primer punto'
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4.6), Inches(3), Inches(1.2))
    sh.fill.solid(); sh.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1; sh.fill.fore_color.brightness = -0.25
    sh.line.fill.background()
    tb = s.shapes.add_textbox(Inches(5), Inches(4.6), Inches(4), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = 'Acento 2, más claro'
    r.font.size = Pt(24); r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_2; r.font.color.brightness = 0.4


patch(build(office), 'office.pptx')

# ---- noche.pptx ---------------------------------------------------------------
NOCHE = ['10172A', 'FFFFFF', '1E293B', 'E2E8F0', 'F59E0B', '38BDF8', 'A3E635', 'F472B6', 'C084FC', '2DD4BF', '38BDF8', 'C084FC']
SECTION_BG = '<p:bg><p:bgPr><a:solidFill><a:schemeClr val="accent1"><a:lumMod val="50000"/></a:schemeClr></a:solidFill><a:effectLst/></p:bgPr></p:bg>'


def noche(prs):
    for i, (lay, t) in enumerate([(0, 'Noche'), (2, 'Sección'), (5, 'Solo el título, claro')]):
        s = prs.slides.add_slide(prs.slide_layouts[lay]); s.shapes.title.text = t
        if lay == 0: s.placeholders[1].text = 'Un tema oscuro'


patch(build(noche), 'noche.pptx', {
    'ppt/theme/theme1.xml': lambda x: retheme(x, 'Noche', NOCHE, 'Cambria', 'Segoe UI'),
    # Dark, as Office's dark themes: the colour map swaps them (background = dk1, text = lt1).
    'ppt/slideMasters/slideMaster1.xml': lambda x: x.replace('<p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef>', '<p:bgRef idx="1003"><a:schemeClr val="bg1"/></p:bgRef>')
        .replace('bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2"', 'bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2"'),
    'ppt/slideLayouts/slideLayout3.xml': lambda x: x.replace('<p:cSld name="Section Header">', '<p:cSld name="Section Header">' + SECTION_BG),
    'ppt/slideLayouts/slideLayout6.xml': lambda x: x.replace('<p:cSld name="Title Only">', '<p:cSld name="Title Only"><p:bg><p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef></p:bg>')
        .replace('<a:masterClrMapping/>', '<a:overrideClrMapping bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'),
})

# ---- faceta.potx (no slides) --------------------------------------------------
FACETA = ['000000', 'FFFFFF', '2C3C43', 'EBEBEB', '90C226', '54A021', 'E6B91E', 'E76618', 'C42F1A', '918655', '99CA3C', 'A9D16F']
patch(build(lambda p: None), 'faceta.potx', {
    'ppt/theme/theme1.xml': lambda x: retheme(x, 'Faceta', FACETA, 'Trebuchet MS', 'Trebuchet MS'),
    '[Content_Types].xml': lambda x: x.replace('presentationml.presentation.main+xml', 'presentationml.template.main+xml'),
})

# ---- brisa.thmx (a theme alone) -----------------------------------------------
BRISA = ['000000', 'FFFFFF', '1B3A4B', 'E8F1F2', '006D77', '83C5BE', 'E29578', 'FFDDD2', '3D5A80', 'EE6C4D', '006D77', '3D5A80']
theme = zipfile.ZipFile(os.path.join(HERE, 'office.pptx')).read('ppt/theme/theme1.xml').decode('utf-8')
with zipfile.ZipFile(os.path.join(HERE, 'brisa.thmx'), 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
               '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
               '<Override PartName="/theme/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/></Types>')
    z.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
               '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="theme/theme/theme1.xml"/></Relationships>')
    z.writestr('theme/theme/theme1.xml', retheme(theme, 'Brisa', BRISA, 'Century Gothic', 'Century Gothic'))

# ---- google.pptx --------------------------------------------------------------
STREAMLINE = ['000000', 'FFFFFF', '3F3F3F', 'F3F3F3', 'FF5E3A', 'FF9B57', '2D8BBA', '41B6E6', '6AA84F', 'F1C232', '2D8BBA', '674EA7']


def google(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = 'Desde Google Slides'; s.placeholders[1].text = 'Tema Streamline'
    s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = 'Título y cuerpo'; s.placeholders[1].text = 'Texto'


names = {1: 'TITLE', 2: 'TITLE_AND_BODY', 3: 'SECTION_HEADER', 6: 'TITLE_ONLY', 7: 'BLANK'}
patch(build(google), 'google.pptx', {
    'ppt/theme/theme1.xml': lambda x: retheme(x, 'Streamline', STREAMLINE, 'Montserrat', 'Lato'),
    **{f'ppt/slideLayouts/slideLayout{i}.xml': (lambda n: lambda x: re.sub(r'<p:cSld name="[^"]*"', f'<p:cSld name="{n}"', x))(n) for i, n in names.items()},
})

# ---- LibreOffice ----------------------------------------------------------------
TPL = '/usr/lib/libreoffice/share/template/common/presnt/Midnightblue.otp'
if shutil.which('soffice') and os.path.exists(TPL):
    with tempfile.TemporaryDirectory() as tmp:
        for fmt in ('pptx', 'odp'):
            subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', tmp, TPL], check=True, capture_output=True)
            shutil.copy(os.path.join(tmp, 'Midnightblue.' + fmt), os.path.join(HERE, 'midnightblue.' + fmt))
print('ok')
